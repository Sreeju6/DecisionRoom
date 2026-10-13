"""Unit tests for multi-format document parser."""
import io
import pytest
import pandas as pd
import fitz  # PyMuPDF
import docx
import openpyxl

from tools.document_parser import (
    DocumentParser,
    FileValidationError,
    ParsingError,
    compute_sha256,
)


def create_sample_pdf() -> bytes:
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "DecisionRoom Technical Architecture. Page 1 content on microservices.")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "DecisionRoom Financial Projections. Page 2 content on revenue and margins.")
    out = doc.tobytes()
    doc.close()
    return out


def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("The company achieved 40% year over year revenue growth.")
    doc.add_heading("Risks", level=2)
    doc.add_paragraph("Customer acquisition costs increased in Q3.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "CAC"
    table.rows[1].cells[1].text = "$120"
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def create_sample_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Financials"
    ws1.append(["Quarter", "Revenue", "GrossMargin"])
    ws1.append(["Q1", 100000, 0.65])
    ws1.append(["Q2", 150000, 0.68])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_parse_pdf():
    pdf_bytes = create_sample_pdf()
    chunks = DocumentParser.parse("architecture.pdf", pdf_bytes, company_id="acme", project_id="p1")
    assert len(chunks) >= 2
    assert any("Page 1" in c.page_or_section for c in chunks)
    assert any("Page 2" in c.page_or_section for c in chunks)
    assert any("microservices" in c.text for c in chunks)
    assert chunks[0].company_id == "acme"


def test_parse_docx():
    docx_bytes = create_sample_docx()
    chunks = DocumentParser.parse("summary.docx", docx_bytes, company_id="acme", project_id="p1")
    assert len(chunks) >= 2
    assert any("Executive Summary" in c.page_or_section for c in chunks)
    assert any("revenue growth" in c.text for c in chunks)
    assert any("Table 1" in c.page_or_section for c in chunks)


def test_parse_txt():
    txt_bytes = "Quarterly Business Review.\nOperations are stable.\nCustomer churn is low.".encode("utf-8")
    chunks = DocumentParser.parse("qbr.txt", txt_bytes, company_id="acme", project_id="p1")
    assert len(chunks) >= 1
    assert "Customer churn" in chunks[0].text


def test_parse_csv():
    df = pd.DataFrame({"Item": ["Server", "Storage"], "Cost": [5000, 2000]})
    csv_bytes = df.to_csv(index=False).encode("utf-8")
    chunks = DocumentParser.parse("costs.csv", csv_bytes, company_id="acme", project_id="p1")
    assert len(chunks) >= 2  # summary + row batch
    assert any("Server" in c.text for c in chunks)


def test_parse_xlsx():
    xlsx_bytes = create_sample_xlsx()
    chunks = DocumentParser.parse("model.xlsx", xlsx_bytes, company_id="acme", project_id="p1")
    assert len(chunks) >= 2
    assert any("Financials" in c.page_or_section for c in chunks)
    assert any("GrossMargin" in c.text for c in chunks)


def test_file_validation_empty():
    with pytest.raises(FileValidationError):
        DocumentParser.parse("empty.txt", b"")


def test_file_validation_unsupported_extension():
    with pytest.raises(FileValidationError):
        DocumentParser.parse("script.exe", b"binary content")


def test_file_validation_corrupt_pdf():
    with pytest.raises(ParsingError):
        DocumentParser.parse("corrupt.pdf", b"not a real pdf content")
